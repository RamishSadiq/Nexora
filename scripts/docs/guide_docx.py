"""Small reusable Markdown to Word renderer for the beginner documentation."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

INLINE = re.compile(r'(\[[^\]]+\]\(<[^>]+>\)|\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`)')

def link(p, label, target):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    el = OxmlElement('w:hyperlink')
    if target.startswith('#'):
        el.set(qn('w:anchor'), target[1:])
    else:
        el.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    colour = OxmlElement('w:color'); colour.set(qn('w:val'), '263D55')
    props.append(colour); run.append(props)
    text = OxmlElement('w:t'); text.text = label
    run.append(text); el.append(run); p._p.append(el)

def inline(p, value):
    for part in INLINE.split(value):
        if not part: continue
        if part.startswith('['):
            m = re.match(r'\[([^\]]+)\]\((?:<([^>]+)>|([^)]*))\)', part)
            if m: link(p, m[1], m[2] or m[3]); continue
        r = p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'): r.bold = True
        if part.startswith('`'): r.font.name = 'Consolas'; r.font.size = Pt(9.5)

def bookmark(p, name, index):
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(index)); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(index))
    p._p.insert(0,start); p._p.append(end)

def font(style, name, size, bold=False):
    style.font.name = name; style.font.size = Pt(size); style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rf = OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+attr),name)
    style.element.get_or_add_rPr().append(rf)

def make_table(doc, rows):
    header = rows[0]; data = [header] + rows[2:]
    n = len(header)
    if n==2: proportions=[0.43,0.57] if 'What you see' in header[0] else [0.30,0.70]
    elif n==3:
        key=header[0].strip()
        proportions={'Role':[0.19,0.46,0.35], 'Request':[0.34,0.42,0.24],
                     'Route or layout scope':[0.47,0.33,0.20],
                     'Context':[0.46,0.30,0.24], 'Module':[0.21,0.59,0.20]}.get(key,[0.38,0.35,0.27])
    else: proportions=[1/n]*n
    t = doc.add_table(rows=0, cols=n); t.autofit=False
    for col,r in zip(t.columns,proportions): col.width=Inches(7*r)
    pr = t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for index,cells in enumerate(data):
        row=t.add_row(); row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index==0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,text in enumerate(cells):
            cell=row.cells[j]; cell.width=Inches(7*proportions[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.style='Table Body'; inline(p,text.strip())
            if index==0:
                p.paragraph_format.keep_with_next=True
                for r in p.runs:r.bold=True
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF3'); cell._tc.get_or_add_tcPr().append(shade)
            margins=OxmlElement('w:tcMar')
            for edge,size in [('top',70),('bottom',70),('left',85),('right',85)]:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),str(size)); e.set(qn('w:type'),'dxa'); margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def build(markdown, destination, image_base, title):
    doc=Document(); s=doc.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=Inches(.65); s.bottom_margin=Inches(.65); s.left_margin=Inches(.75); s.right_margin=Inches(.75)
    s.footer_distance=Inches(.25)
    font(doc.styles['Normal'],'Arial',11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(4); doc.styles['Normal'].paragraph_format.line_spacing=1.08
    for name,size in [('Title',28),('Subtitle',13),('Heading 1',18),('Heading 2',13),('Heading 3',11.5),('Heading 4',11)]:
        font(doc.styles[name],'Arial',size,name.startswith('Heading'))
        pf=doc.styles[name].paragraph_format; pf.space_before=Pt(12 if name.startswith('Heading') else 0); pf.space_after=Pt(7); pf.keep_with_next=True
    for name,size,face in [('Guide Code',9.5,'Consolas'),('Table Body',10,'Arial'),('Guide Small',9.5,'Arial')]:
        doc.styles.add_style(name,1); font(doc.styles[name],face,size)
    cp=doc.styles['Guide Code'].paragraph_format
    cp.space_after=Pt(0); cp.space_before=Pt(0); cp.line_spacing=1.02; cp.keep_together=True
    cp.left_indent=Pt(32); cp.first_line_indent=Pt(-32)
    doc.styles['Table Body'].paragraph_format.space_after=Pt(1); doc.styles['Table Body'].paragraph_format.line_spacing=1.04
    f=s.footer.paragraphs[0]; f.style='Guide Small'; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('Nexora learning guide    '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
    doc.core_properties.title=title; doc.core_properties.author='Nexora'; doc.core_properties.subject='Beginner guidance with verified source maps'
    lines=markdown.splitlines(); i=0; in_code=False; code=[]; bkm=1
    while i<len(lines):
        value=lines[i]
        if value.startswith('<!--'):
            if value.startswith('<!-- CHAPTER '):
                p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
            i+=1; continue
        if value.startswith('```'):
            in_code=not in_code
            if not in_code:
                for p in code[-3:-1]:p.paragraph_format.keep_with_next=True
                if code:code[-1].paragraph_format.keep_with_next=False
                code=[]; doc.add_paragraph().paragraph_format.space_after=Pt(3)
            i+=1; continue
        if in_code:
            p=doc.add_paragraph(value or ' ','Guide Code')
            if len(code)<2:p.paragraph_format.keep_with_next=True
            code.append(p); i+=1; continue
        if not value.strip():i+=1;continue
        if value.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                rows.append(lines[i].strip().strip('|').split('|')); i+=1
            make_table(doc,rows);continue
        heading=re.match(r'^(#{1,5}) (.*)$',value)
        if heading:
            depth=len(heading[1]); text=heading[2]
            p=doc.add_paragraph(text,'Title' if depth==1 else 'Heading '+str(min(depth-1,4)))
            m=re.match(r'(\d{2}) Nexora',text)
            if m:bookmark(p,'guide_'+m[1],bkm); bkm+=1
        elif value.startswith('!['):
            m=re.match(r'!\[([^]]+)\]\(([^)]+)\)',value)
            if not m:raise ValueError(value)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            picture=p.add_run().add_picture(str(Path(image_base)/m[2]),width=Inches(6.95));picture._inline.docPr.set('descr',m[1])
            p.paragraph_format.keep_with_next=True
        else:
            p=doc.add_paragraph(style='Guide Small' if value.startswith('Source reference ') else 'Normal')
            if value.startswith('- '):value='• '+value[2:];p.paragraph_format.left_indent=Inches(.16);p.paragraph_format.first_line_indent=Inches(-.16)
            inline(p,value.rstrip())
            if value.endswith('  '):p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
            if value.startswith('Source reference '):p.paragraph_format.keep_with_next=True
        i+=1
    for style in doc.styles:
        for e in style.element.findall('.//'+qn('w:pBdr')):e.getparent().remove(e)
    for p in doc.paragraphs:
        for e in p._p.findall('.//'+qn('w:pBdr')):e.getparent().remove(e)
    doc.save(destination)
    return doc
